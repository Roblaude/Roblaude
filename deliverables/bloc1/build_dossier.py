#!/usr/bin/env python3
"""Construit le dossier individuel de conception B1 de RobLaude."""

from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
DIAGRAMS = HERE / "diagrams"
OUTPUT = HERE / "Dossier_Conception_B1_RobLaude_Wissem.docx"

NAVY = "15384A"
TEAL = "1F6F78"
CORAL = "E46F61"
INK = "17313B"
MUTED = "5F7077"
PALE = "EAF3F3"
PALE_BLUE = "EAF0F5"
PALE_CORAL = "FCEDEA"
LIGHT = "F7F9F9"
WHITE = "FFFFFF"
GREEN = "3C8067"
RED = "A34F4F"
PURPLE = "6B5870"
FONT_NAME = "Arial"


def rgb(value):
    return RGBColor.from_string(value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table, widths_cm):
    table.autofit = False
    total_dxa = int(sum(widths_cm) / 2.54 * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_cm in widths_cm:
        width_dxa = int(width_cm / 2.54 * 1440)
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width_dxa))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width_dxa = int(widths_cm[idx] / 2.54 * 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_dxa))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_run(run, size=None, color=INK, bold=None, italic=None, font_name=FONT_NAME):
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run(run, size=9, color=MUTED)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, color, before, after in (
        ("Heading 1", 19, NAVY, 12, 9),
        ("Heading 2", 13.5, TEAL, 10, 6),
        ("Heading 3", 11.5, NAVY, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = FONT_NAME
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = rgb(MUTED)
    caption.font.italic = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT_NAME
        style.font.size = Pt(10.5)
        style.font.color.rgb = rgb(INK)
        style.paragraph_format.left_indent = Cm(0.7)
        style.paragraph_format.first_line_indent = Cm(-0.35)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.2

    if "Source" not in styles:
        source_style = styles.add_style("Source", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source_style = styles["Source"]
    source_style.font.name = FONT_NAME
    source_style.font.size = Pt(8.5)
    source_style.font.color.rgb = rgb(MUTED)
    source_style.paragraph_format.space_before = Pt(2)
    source_style.paragraph_format.space_after = Pt(6)
    source_style.paragraph_format.line_spacing = 1.05

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("ROBLAUDE  |  DOSSIER DE CONCEPTION B1")
    set_run(r, size=8.5, color=TEAL, bold=True)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D3E0E3")
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run("HETIC Web3 2026   •   ")
    set_run(run, size=8.5, color=MUTED)
    add_page_field(fp)
    return doc


def add_para(doc, text="", *, bold_lead=None, align=None, color=None,
             size=None, italic=False, after=None, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run(lead, size=size, color=color or INK, bold=True, italic=italic)
        rest = p.add_run(text[len(bold_lead):])
        set_run(rest, size=size, color=color or INK, italic=italic)
    else:
        run = p.add_run(text)
        set_run(run, size=size, color=color or INK, italic=italic)
    return p


def add_source(doc, text):
    p = doc.add_paragraph(style="Source")
    p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_callout(doc, label, text, fill=PALE, accent=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    set_table_width(table, [16.4])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    set_run(r, size=9, color=accent, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(text)
    set_run(r2, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_table(doc, headers, rows, widths, header_fill=NAVY, small=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run(run, size=8.5 if small else 9, color=WHITE, bold=True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            run = p.add_run(str(value))
            set_run(run, size=8 if small else 8.7, color=INK)
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def set_picture_alt(inline_shape, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def add_figure(doc, filename, caption, width_cm=16.2, alt=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    shape = run.add_picture(str(DIAGRAMS / filename), width=Cm(width_cm))
    set_picture_alt(shape, alt or caption)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.paragraph_format.keep_with_next = False


def page_break(doc):
    # Les chapitres gèrent leur propre rupture. Une rupture isolée peut créer
    # une page blanche lorsque le contenu précédent déborde déjà naturellement.
    return None


def heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    if level == 1:
        paragraph.paragraph_format.page_break_before = True
    return paragraph


def cover(doc):
    for _ in range(4):
        add_para(doc, "", after=4)
    p = add_para(doc, "HETIC  •  WEB3  •  PROMOTION 2026", align=WD_ALIGN_PARAGRAPH.CENTER,
                 color=CORAL, size=10, after=22)
    p.runs[0].bold = True

    title = add_para(doc, "DOSSIER DE CONCEPTION", align=WD_ALIGN_PARAGRAPH.CENTER,
                     color=NAVY, size=27, after=4)
    title.runs[0].bold = True
    subtitle = add_para(doc, "RobLaude", align=WD_ALIGN_PARAGRAPH.CENTER,
                        color=TEAL, size=34, after=8)
    subtitle.runs[0].bold = True
    add_para(doc, "Robot d'assistance autonome pour les personnes à mobilité réduite en ERP",
             align=WD_ALIGN_PARAGRAPH.CENTER, color=MUTED, size=14, after=48)

    add_para(doc, "ÉPREUVE 1  •  BLOC 1 : CONCEPTION", align=WD_ALIGN_PARAGRAPH.CENTER,
             color=NAVY, size=12, after=22).runs[0].bold = True

    add_callout(
        doc,
        "Objet du dossier",
        "Analyser le besoin, expliquer les arbitrages et formaliser une architecture cohérente, en distinguant la cible initiale, la réalisation observée et les améliorations encore nécessaires.",
        fill=PALE_BLUE,
        accent=NAVY,
    )
    add_para(doc, "Travail individuel : Wissem Karboub", align=WD_ALIGN_PARAGRAPH.CENTER,
             color=INK, size=11, after=3).runs[0].bold = True
    add_para(doc, "Titre RNCP 36146  •  Concepteur Développeur de Solutions Digitales",
             align=WD_ALIGN_PARAGRAPH.CENTER, color=MUTED, size=9.5, italic=True, after=2)
    add_para(doc, "Année 2026", align=WD_ALIGN_PARAGRAPH.CENTER, color=MUTED, size=9.5)


def contents_and_summary(doc):
    doc.add_page_break()
    for _ in range(4):
        add_para(doc, "", after=4)
    toc_title = add_para(doc, "SOMMAIRE", align=WD_ALIGN_PARAGRAPH.CENTER,
                         color=NAVY, size=25, after=24)
    toc_title.runs[0].bold = True
    toc = [
        ("1. Résumé exécutif", 3),
        ("2. Contexte et problématique", 4),
        ("3. Parties prenantes, besoins et priorités", 5),
        ("4. Périmètre fonctionnel", 6),
        ("5. État de l'art et arbitrages", 7),
        ("6. Architecture de la solution", 9),
        ("7. Contrats de communication et données", 11),
        ("8. Spécifications fonctionnelles", 13),
        ("9. Réalisation et validation", 16),
        ("10. Contraintes, risques et limites", 17),
        ("11. Axes d'amélioration et évolutivité", 18),
        ("12. Traçabilité des besoins", 19),
        ("13. Conclusion et sources", 20),
    ]
    for item, page in toc:
        p = add_para(doc, "", color=INK, size=12.5, after=7)
        p.paragraph_format.left_indent = Cm(2.1)
        p.paragraph_format.right_indent = Cm(2.1)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        item_run = p.add_run(item)
        set_run(item_run, size=12.5, color=INK)
        p.add_run("\t")
        page_run = p.add_run(str(page))
        set_run(page_run, size=12.5, color=TEAL, bold=True)

    doc.add_page_break()
    heading(doc, "1. Résumé exécutif", 1)
    add_para(doc, "RobLaude est un prototype de robot d'assistance destiné aux personnes à mobilité réduite dans les établissements recevant du public. Le projet répond à une situation simple : dans un bâtiment administratif, un usager peut dépendre d'un tiers pour déplacer un document ou récupérer un objet. La solution imaginée associe une interface web accessible, un serveur chargé des missions et un robot mobile capable de se repérer, de naviguer et de remonter son état.")
    add_para(doc, "L'architecture sépare les responsabilités. L'interface reste centrée sur l'utilisateur. Le backend contrôle les comptes, les points et le cycle des missions. MQTT transporte les commandes et la télémétrie entre le web et le robot. ROS 2 coordonne les capteurs, la navigation et l'actionnement. Les composants peuvent ainsi évoluer sans créer de dépendance directe entre l'application et le matériel.")
    add_para(doc, "La navigation, la cartographie et plusieurs flux de communication ont été testés sur le robot réel. La récupération autonome d'un objet reste plus fragile : le bras et la pince répondent, mais la chaîne complète de perception, saisie et transport demande encore de la stabilisation. Cette fonction est donc présentée comme partiellement validée.")
    add_callout(doc, "Recommandation", "Conserver l'architecture en couches et concentrer la prochaine phase sur la sûreté, la robustesse réseau, la validation de la saisie et la réduction de la charge embarquée.", fill=PALE, accent=TEAL)

    add_table(doc,
              ["Compétence", "Ce que le dossier démontre"],
              [
                  ("C1.1", "Besoin, acteurs, objectifs hiérarchisés, contraintes et évolutions probables."),
                  ("C1.2", "Comparaisons ciblées, critères de décision et conséquences des arbitrages."),
                  ("C1.3", "Architecture, contrats, spécifications et diagrammes cohérents avec le projet."),
              ], [2.5, 13.9])


def context_and_needs(doc):
    page_break(doc)
    heading(doc, "2. Contexte et problématique", 1)
    heading(doc, "2.1 Une difficulté d'usage dans les ERP", 2)
    add_para(doc, "Les mairies, hôpitaux et bâtiments administratifs regroupent des services dispersés dans des couloirs, sur plusieurs zones ou à différents étages. Pour une personne à mobilité réduite, un déplacement supplémentaire pour transmettre un document ou récupérer un objet peut représenter un effort important. L'alternative consiste souvent à solliciter un agent ou un accompagnant. Le problème ne se résume donc pas au transport : il touche directement l'autonomie de l'usager et l'organisation du personnel.")
    add_para(doc, "RobLaude a été conçu comme un service de proximité dans le bâtiment. L'utilisateur formule une demande depuis son téléphone ou un poste web. Le robot rejoint un point connu, transporte le document ou tente de récupérer un objet, puis se rend à la destination. Pendant le trajet, l'utilisateur doit comprendre ce qui se passe et pouvoir arrêter le mouvement.")

    heading(doc, "2.2 Enjeux métier", 2)
    add_table(doc,
              ["Enjeu", "Traduction pour RobLaude", "Conséquence de conception"],
              [
                  ("Autonomie", "Limiter le recours systématique à un tiers.", "Parcours guidé et commandes accessibles."),
                  ("Confiance", "Rendre l'état de la mission compréhensible.", "Suivi temps réel et statuts explicites."),
                  ("Sécurité", "Permettre une interruption immédiate.", "Arrêt dédié et gestion de la pause."),
                  ("Continuité", "Éviter qu'une panne d'un écran bloque tout le système.", "Séparation frontend, backend et robot."),
                  ("Maintenabilité", "Faire évoluer le prototype avec une petite équipe.", "Standards existants et composants isolés."),
              ], [2.5, 6.0, 7.9], small=True)

    heading(doc, "2.3 Retombées attendues", 2)
    add_para(doc, "Les retombées sont formulées comme des objectifs, car le projet académique n'a pas conduit d'étude économique en exploitation. La solution vise à réduire certains déplacements d'assistance, à fluidifier les échanges internes et à rendre le service plus inclusif. Elle pourrait également offrir au personnel technique une vision centralisée des missions et de l'état du robot.")
    add_callout(doc, "Limite de l'analyse", "Aucun gain financier chiffré n'est annoncé. Il faudrait une expérimentation en ERP, un volume réel de missions et des coûts de maintenance observés pour calculer un retour sur investissement crédible.", fill=PALE_CORAL, accent=CORAL)
    add_source(doc, "Sources projet : [R1] Cahier des charges RobLaude v1.2 ; [R2] Architecture canonique.")

    page_break(doc)
    heading(doc, "3. Parties prenantes, besoins et priorités", 1)
    heading(doc, "3.1 Acteurs", 2)
    add_table(doc,
              ["Acteur", "Attente principale", "Point de vigilance"],
              [
                  ("Utilisateur PMR", "Créer et suivre une mission sans assistance technique.", "Lisibilité, navigation clavier et action STOP visible."),
                  ("Agent d'accueil", "Utiliser le service et accompagner un chargement manuel.", "Compréhension immédiate des statuts."),
                  ("Administrateur", "Configurer points, objets et consulter l'historique.", "Droits distincts et traçabilité."),
                  ("Personnel technique", "Diagnostiquer le robot et maintenir la carte.", "Accès aux états sans exposer ROS 2 au public."),
                  ("Robot", "Exécuter une commande valide et remonter son résultat.", "Ressources de calcul et réseau limités."),
              ], [3.0, 7.3, 6.1], small=True)

    heading(doc, "3.2 Hiérarchisation", 2)
    add_para(doc, "Les fonctions n'ont pas le même poids. La sécurité et la capacité à arrêter le robot passent avant la richesse du suivi. La navigation et le transport de document forment le premier service complet. La récupération d'objet apporte davantage d'autonomie, mais elle dépend d'une chaîne plus risquée : caméra, détection, calibration, bras, pince et validation de la prise.")
    add_table(doc,
              ["Priorité", "Besoins"],
              [
                  ("Critique", "Arrêt d'urgence, refus d'une mission impossible, état du robot connu."),
                  ("Haute", "Créer une mission de transport, naviguer, suivre l'avancement."),
                  ("Moyenne", "Configurer les points, objets et cartes du bâtiment."),
                  ("Évolutive", "Récupération autonome fiable, multi-robot, déploiement multi-site."),
              ], [3.0, 13.4])

    heading(doc, "3.3 Contraintes initiales", 2)
    add_bullets(doc, [
        "Une équipe de deux personnes et dix-huit semaines, avec une expérience plus forte sur le web que sur ROS 2.",
        "Un robot réel équipé d'un Jetson Nano à quatre cœurs et 4 Go de mémoire, ce qui limite les traitements simultanés.",
        "Un réseau local dont les adresses peuvent changer, avec un robot qui doit rester joignable sans configuration manuelle permanente.",
        "Des capteurs et un bras dont les limites devaient être découvertes par des essais sur le matériel.",
        "Une interface destinée à des usages mobiles et à des personnes pouvant naviguer au clavier ou avec des dispositifs adaptés.",
    ])
    add_source(doc, "Sources projet : [R1] Cahier des charges ; [R2] Architecture ; [R3] Roadmap.")


def functional_scope(doc):
    page_break(doc)
    heading(doc, "4. Périmètre fonctionnel", 1)
    add_para(doc, "Le périmètre s'organise autour de sept cas d'utilisation. Quatre concernent directement l'utilisateur : transporter un document, récupérer un objet, suivre une mission et arrêter le robot. Trois relèvent de l'administration : configurer les points, définir les objets saisissables et consulter l'historique.")
    add_figure(doc, "use-cases.png", "Figure 1. Cas d'utilisation retenus pour le dossier de conception.", width_cm=16.4,
               alt="Diagramme des cas d'utilisation de RobLaude avec utilisateur PMR et administrateur")
    heading(doc, "4.1 Périmètre minimal utilisable", 2)
    add_para(doc, "Le premier scénario cohérent est le transport de document. Il mobilise déjà l'ensemble de la chaîne sans dépendre du bras : création de mission, attribution au robot, navigation vers le point de collecte, confirmation du chargement, navigation vers la destination et notification finale. Ce choix réduit le risque tout en validant l'architecture de bout en bout.")
    heading(doc, "4.2 Extension vers la manipulation", 2)
    add_para(doc, "La récupération d'objet reprend le même socle, puis ajoute la détection et la saisie. Elle n'est pas traitée comme un simple bouton supplémentaire. Elle impose des préconditions propres : objet connu, caméra disponible, bras calibré et stratégie de reprise en cas d'échec. Cette séparation évite qu'une fonction expérimentale fragilise le transport simple.")


def state_of_art(doc):
    page_break(doc)
    heading(doc, "5. État de l'art et arbitrages", 1)
    add_para(doc, "L'état de l'art a été limité aux décisions qui modifient réellement le coût, le délai, la performance ou la pérennité de la solution. Le but n'était pas de comparer toutes les bibliothèques disponibles, mais de choisir un chemin compatible avec l'équipe, le matériel et le besoin.")

    heading(doc, "5.1 Interface : PWA ou application native", 2)
    add_table(doc,
              ["Critère", "PWA", "Application native"],
              [
                  ("Accès", "URL puis installation possible.", "Installation via un écosystème mobile."),
                  ("Développement", "Une base React/TypeScript pour plusieurs écrans.", "Développements et validations spécifiques par plateforme."),
                  ("Matériel mobile", "Accès suffisant pour le suivi et les formulaires.", "Meilleure intégration aux API profondes du téléphone."),
                  ("Choix RobLaude", "Retenue : délai court, compétences disponibles, usage web central.", "Non retenue : coût d'apprentissage sans bénéfice décisif ici."),
              ], [2.5, 6.95, 6.95], small=True)
    add_para(doc, "La PWA ne signifie pas que le robot fonctionne hors ligne. Une mission nécessite le réseau. Le service worker sert surtout à l'installation et au chargement de l'interface. Cette limite est cohérente avec le besoin : masquer une perte de connexion pendant une commande robot serait trompeur.")
    add_source(doc, "Référence externe : [E1] MDN, Making PWAs installable. Source projet : [R1], section interface web.")

    heading(doc, "5.2 Communication web : REST et WebSocket", 2)
    add_para(doc, "REST convient aux opérations ponctuelles qui ont une réponse claire : créer une mission, consulter des points ou confirmer un chargement. Le suivi du robot produit au contraire des changements réguliers. Interroger le serveur en boucle augmenterait le trafic et le délai d'affichage. WebSocket permet au backend de pousser une position ou un statut dès qu'il le reçoit. Les deux mécanismes sont donc complémentaires.")

    heading(doc, "5.3 Communication robot : MQTT ou exposition directe de ROS 2", 2)
    add_table(doc,
              ["Critère", "MQTT entre web et robot", "Accès direct à ROS 2"],
              [
                  ("Découplage", "Contrat stable entre deux environnements.", "Le web dépend du graphe ROS et de ses types."),
                  ("Réseau", "Broker, reconnexion, QoS et état retained.", "Découverte DDS plus sensible à la topologie réseau."),
                  ("Sécurité", "Un point de contrôle côté backend et broker.", "Surface ROS 2 exposée au système web."),
                  ("Évolutivité", "Préfixe par robot et abonnements génériques.", "Nommage et découverte à adapter pour chaque robot."),
                  ("Choix", "Retenu pour isoler le domaine web du domaine robotique.", "Réservé aux communications internes du robot."),
              ], [2.4, 7.0, 7.0], small=True)
    add_source(doc, "Références : [E2] spécification OASIS MQTT ; [E3] documentation ROS 2 sur topics, services et actions.")

    page_break(doc)
    heading(doc, "5.4 Framework robotique", 2)
    add_para(doc, "Développer directement les drivers, la navigation, la localisation et les outils de diagnostic aurait dépassé le calendrier. ROS 2 Humble fournit un modèle par nœuds, des messages typés, des topics pour les flux continus et des actions pour les tâches longues. Nav2 expose notamment une action de navigation avec objectif, feedback et résultat, adaptée à un déplacement qui peut être annulé.")
    add_para(doc, "Nav2 et SLAM Toolbox ont donc été retenus plutôt qu'un moteur de navigation écrit sur mesure. Ce choix accélère le prototype et s'appuie sur un écosystème documenté. En contrepartie, il impose de comprendre les repères TF, les horloges, les QoS et le cycle de vie des nœuds. Le problème temporel des LiDAR a montré que cette complexité est réelle.")
    add_source(doc, "Références : [E3] ROS 2 Interfaces ; [E4] documentation Nav2 ; [R2] architecture RobLaude.")

    heading(doc, "5.5 Hébergement des traitements", 2)
    add_table(doc,
              ["Option", "Avantage", "Limite", "Décision"],
              [
                  ("Tout sur le robot", "Autonomie locale et faible dépendance au serveur.", "Charge CPU et mémoire élevée sur le Jetson Nano.", "Conserver la navigation critique ; limiter les flux et traitements non essentiels."),
                  ("Tout sur un serveur", "Plus de ressources de calcul.", "Dépendance réseau trop forte pour les mouvements et capteurs.", "Écartée pour le contrôle robot."),
                  ("Architecture hybride", "Métier côté backend, contrôle temps réel côté robot.", "Contrats et synchronisation à maintenir.", "Retenue."),
              ], [2.8, 4.4, 5.2, 4.0], small=True)

    heading(doc, "5.6 Synthèse des recommandations", 2)
    add_table(doc,
              ["Décision", "Coût et délai", "Performance", "Pérennité"],
              [
                  ("PWA React", "Une base pour mobile et desktop.", "Interface légère ; dépend du réseau pour les missions.", "Technologies web standard et source éditable."),
                  ("MQTT", "Ajoute un broker mais évite un couplage sur mesure.", "QoS adapté à chaque flux ; télémétrie à limiter.", "Contrats versionnés et extension multi-robot."),
                  ("ROS 2 + Nav2", "Courbe d'apprentissage, mais briques déjà disponibles.", "Exécution locale ; réglages nécessaires sur Jetson.", "Composants remplaçables derrière des interfaces ROS."),
                  ("Docker sur Jetson", "Réutilise le matériel livré malgré son OS d'origine.", "Surcoût modéré en ressources.", "Environnement reproductible et dépendances isolées."),
              ], [3.4, 4.3, 4.3, 4.4], small=True)


def architecture(doc):
    page_break(doc)
    heading(doc, "6. Architecture de la solution", 1)
    add_para(doc, "L'architecture repose sur quatre zones : l'expérience utilisateur, les services web, la messagerie et le robot. Une zone peut évoluer tant que son contrat avec la suivante reste stable. Ce principe répond directement à la taille de l'équipe et au caractère expérimental du matériel.")
    add_figure(doc, "components.png", "Figure 2. Architecture logique et responsabilités principales.", width_cm=16.4,
               alt="Architecture logique RobLaude de la PWA au robot en passant par le backend et MQTT")

    heading(doc, "6.1 Frontend", 2)
    add_para(doc, "La PWA présente les missions, les points et l'état du robot. Elle ne décide pas si une mission est possible et ne publie jamais sur MQTT. Cette limite volontaire empêche qu'une modification d'interface contourne les règles métier ou adresse directement des composants physiques.")
    heading(doc, "6.2 Backend et base de données", 2)
    add_para(doc, "Le backend Express authentifie les utilisateurs, valide les demandes et conserve l'état métier dans MySQL via Prisma. Il vérifie notamment la disponibilité du robot et les paramètres de mission avant de publier une commande. Il traduit ensuite les retours MQTT en mises à jour de données et en événements WebSocket.")
    heading(doc, "6.3 Broker et bridge", 2)
    add_para(doc, "Mosquitto reçoit les messages MQTT. Le bridge exécuté sur le robot traduit les commandes en messages ou actions ROS 2 et remonte la télémétrie. Le bridge ne remplace ni le backend ni les nœuds de navigation : il adapte deux modèles de communication différents.")
    heading(doc, "6.4 Domaine robot", 2)
    add_para(doc, "ROS 2 organise la partie embarquée en programmes spécialisés. Le nœud de mission pilote l'action NavigateToPose. SLAM Toolbox produit la carte et la relation entre les repères map et odom. Les drivers publient LiDAR, odométrie et batterie. Le STM32 reste responsable de l'actionnement bas niveau des moteurs et des servos.")

    page_break(doc)
    heading(doc, "6.5 Décomposition interne du robot", 2)
    add_table(doc,
              ["Composant", "Responsabilité", "Entrées", "Sorties"],
              [
                  ("mqtt_bridge", "Adapter MQTT et ROS 2.", "Commandes MQTT, topics ROS.", "Topics ROS, télémétrie MQTT."),
                  ("mission_executor", "Orchestrer le scénario de mission.", "Commandes internes, détections.", "Objectifs Nav2, statuts, bras."),
                  ("scan_restamper", "Corriger le temps des scans fusionnés.", "/scan_multi", "/scan_fixed"),
                  ("SLAM Toolbox", "Construire la carte et localiser.", "Scans et TF.", "/map et TF map→odom."),
                  ("Nav2", "Planifier et exécuter un déplacement.", "Goal NavigateToPose, carte, obstacles.", "Commandes de vitesse, feedback, résultat."),
                  ("mapping_supervisor", "Démarrer, arrêter et sauvegarder une cartographie.", "Commandes de mapping.", "État et résultat de sauvegarde."),
                  ("STM32 / micro-ROS", "Piloter le matériel bas niveau.", "Vitesse, bras, pince.", "Odométrie, IMU, batterie."),
              ], [3.1, 4.5, 4.4, 4.4], small=True)
    add_callout(doc, "Pourquoi cette granularité", "Chaque nœud a une responsabilité identifiable. Un problème de timestamp LiDAR peut être corrigé sans modifier le backend, la base de données ou l'interface.", fill=PALE, accent=TEAL)

    heading(doc, "6.6 Déploiement", 2)
    add_para(doc, "Le Jetson Nano exécute ROS 2 Humble dans le conteneur m3pro. Le workspace RobLaude est monté dans ce conteneur et construit avec colcon. Un second conteneur héberge l'agent micro-ROS qui relie le STM32 au graphe ROS 2. Cette solution a permis de conserver le matériel livré tout en utilisant les versions logicielles nécessaires au projet.")
    add_para(doc, "Le backend, MySQL et Mosquitto peuvent être lancés avec Docker Compose sur la machine de développement. Le robot retrouve le broker sur le réseau local. L'adresse du robot étant attribuée par DHCP, un script le localise par son adresse MAC avant les opérations de déploiement.")
    add_source(doc, "Sources projet : [R2] Architecture canonique ; [R4] Spécification MQTT ; scripts de déploiement du dossier robot.")


def contracts_and_data(doc):
    page_break(doc)
    heading(doc, "7. Contrats de communication et données", 1)
    heading(doc, "7.1 Arborescence MQTT", 2)
    add_para(doc, "Chaque robot possède un sous-arbre `roblaude/{robotId}/...`. Cette convention évite de coder un ensemble de topics différent pour chaque machine et prépare la gestion de plusieurs robots. Les messages se répartissent entre commandes, télémétrie, état global et cycle de mission.")
    add_table(doc,
              ["Famille", "Direction", "Exemples", "Règle"],
              [
                  ("cmd", "Backend → Robot", "mission, cancel, resume, emergency-stop", "Événements non retained ; identifiant pour l'idempotence."),
                  ("telemetry", "Robot → Backend", "position, batterie, carte", "États périodiques ; fréquence adaptée à l'usage."),
                  ("status", "Robot → Backend", "AVAILABLE, BUSY, ERROR", "État retained pour une reprise rapide."),
                  ("mission", "Robot → Backend", "ack, status, result", "Accusé, progression et sortie clairement séparés."),
                  ("connection", "Robot → Backend", "online true/false", "Last Will pour signaler une rupture non propre."),
              ], [2.4, 3.2, 5.0, 5.8], small=True)

    heading(doc, "7.2 QoS et nature des messages", 2)
    add_para(doc, "Le niveau de QoS n'est pas identique pour tous les flux. Une position sera remplacée quelques instants plus tard ; la perte ponctuelle d'un échantillon est tolérable. Une commande de mission ou un résultat final exige davantage de garantie. Les états utiles à la reconnexion sont retained, tandis qu'une commande ne l'est jamais afin d'éviter son rejeu involontaire.")
    add_table(doc,
              ["Type", "QoS retenu", "Retained", "Justification"],
              [
                  ("Position", "0", "Oui", "Flux fréquent ; la prochaine valeur remplace la précédente."),
                  ("Batterie / état", "1", "Oui", "Un nouvel abonné doit connaître le dernier état."),
                  ("Commande", "2", "Non", "Éviter perte ou doublon pendant une session, sans rejouer un ordre ancien."),
                  ("Résultat final", "2", "Non", "La clôture d'une mission doit être traitée une seule fois."),
              ], [3.1, 2.3, 2.3, 8.7], small=True)
    add_callout(doc, "Nuance technique", "MQTT QoS 2 ne garantit pas qu'une commande publiée pendant une longue absence du robot sera toujours exécutée. Le backend doit aussi vérifier la présence et refuser une demande devenue incohérente.", fill=PALE_CORAL, accent=CORAL)

    heading(doc, "7.3 Données métier", 2)
    add_para(doc, "Le schéma Prisma matérialise les responsabilités du backend. Un utilisateur crée une mission. Une mission référence des points, un robot et éventuellement un objet. Les statuts représentent le cycle de vie ; ils ne sont pas déduits uniquement de la dernière position reçue. Cette distinction permet de conserver un historique, d'expliquer un échec et d'éviter qu'une coupure réseau efface la demande.")
    add_table(doc,
              ["Entité", "Rôle", "Relations principales"],
              [
                  ("User", "Identité et rôle.", "Crée des missions."),
                  ("Mission", "Demande, type, statut et résultat.", "Utilisateur, robot, points, objet éventuel."),
                  ("Robot", "État opérationnel et identification.", "Exécute plusieurs missions dans le temps."),
                  ("Point", "Emplacement nommé et coordonnées.", "Départ ou destination."),
                  ("GraspObject", "Objet configurable pour la saisie.", "Associé aux missions de récupération."),
                  ("MapSnapshot", "Carte sauvegardée et métadonnées.", "Liée aux sessions de cartographie."),
              ], [3.0, 6.3, 7.1], small=True)
    add_source(doc, "Sources projet : [R4] Spécification MQTT ; `web/backend/prisma/schema.prisma`.")


def specifications(doc):
    page_break(doc)
    heading(doc, "8. Spécifications fonctionnelles", 1)
    heading(doc, "8.1 Mission de transport", 2)
    add_figure(doc, "sequence-uc01.png", "Figure 3. Flux simplifié d'une mission de transport.", width_cm=16.4,
               alt="Diagramme de séquence depuis l'utilisateur jusqu'à Nav2")
    add_table(doc,
              ["Élément", "Spécification"],
              [
                  ("Préconditions", "Utilisateur authentifié ; points configurés ; robot disponible et connecté."),
                  ("Déclencheur", "Validation du départ et de la destination dans la PWA."),
                  ("Scénario nominal", "Création en base, commande MQTT, navigation vers collecte, confirmation du chargement, navigation vers destination, résultat."),
                  ("Suivi", "Position et sous-état transmis au frontend via WebSocket."),
                  ("Erreurs", "Robot hors ligne, refus, timeout, obstacle bloquant, annulation ou arrêt d'urgence."),
                  ("Postcondition", "Mission terminée, échouée ou annulée avec un état et une raison enregistrés."),
              ], [3.3, 13.1], small=True)

    heading(doc, "8.2 Arrêt d'urgence", 2)
    add_para(doc, "Le bouton STOP doit rester accessible quel que soit l'écran. Son activation envoie une requête au backend, qui publie une commande dédiée. Le robot interrompt la navigation et renvoie un état de pause. L'utilisateur peut ensuite reprendre ou annuler. Cette logique évite de confondre une pause de sécurité avec un échec définitif.")
    add_callout(doc, "Exigence de sûreté", "Le dispositif logiciel contribue à l'arrêt, mais ne remplace pas une chaîne d'arrêt d'urgence certifiée au niveau matériel. Une exploitation en ERP demanderait une analyse de risques et des validations réglementaires complémentaires.", fill=PALE_CORAL, accent=RED)

    page_break(doc)
    heading(doc, "8.3 Cycle de vie", 2)
    add_figure(doc, "mission-states.png", "Figure 4. États métier principaux d'une mission.", width_cm=16.4,
               alt="Diagramme d'états d'une mission RobLaude")
    add_para(doc, "Le backend conserve l'état métier de référence. Le robot exécute et publie ses retours, mais une position ou un topic isolé ne suffit pas à redéfinir la mission. Cette règle facilite la reprise après une reconnexion et conserve une trace compréhensible pour l'utilisateur.")

    grasp_heading = heading(doc, "8.4 Récupération d'objet", 2)
    grasp_heading.paragraph_format.page_break_before = True
    add_para(doc, "Une mission de récupération ajoute quatre sous-étapes : rejoindre l'emplacement, détecter l'objet, effectuer la saisie puis transporter et déposer. Les erreurs de détection ou de prise doivent être distinguées d'une panne générale. Lorsque le robot reste opérationnel, l'utilisateur peut repositionner l'objet ou relancer une mission.")
    add_table(doc,
              ["Sous-étape", "Entrée", "Résultat attendu", "Sortie d'erreur"],
              [
                  ("Détection", "Couleur ou objet ciblé, image caméra.", "Position exploitable de l'objet.", "Objet non trouvé."),
                  ("Approche", "Position de l'objet.", "Robot et bras correctement placés.", "Position inaccessible."),
                  ("Saisie", "Pose du bras et commande pince.", "Objet maintenu.", "Prise échouée ou incertaine."),
                  ("Transport", "Objet saisi, destination valide.", "Arrivée sans perte.", "Objet tombé ou navigation impossible."),
                  ("Dépôt", "Destination atteinte.", "Objet libéré et bras au repos.", "Dépôt non confirmé."),
              ], [2.4, 4.3, 4.5, 5.2], small=True)


def validation_and_limits(doc):
    page_break(doc)
    heading(doc, "9. Réalisation et validation", 1)
    add_para(doc, "La présence d'un composant dans le dépôt prouve son implémentation, pas son fonctionnement sur le terrain. La validation ci-dessous sépare donc le code disponible, les essais documentés et les éléments encore fragiles.")
    add_table(doc,
              ["Élément", "État retenu", "Preuve disponible"],
              [
                  ("PWA et parcours de mission", "Implémentés", "Code React, tests frontend et runbook de démonstration."),
                  ("API, données et MQTT", "Implémentés", "Code Express/Prisma, bridge et tests MQTT."),
                  ("Navigation Nav2", "Testée sur robot réel", "Historique de sessions et goals atteints documentés."),
                  ("Cartographie", "Testée sur robot réel", "Carte sauvegardée et procédure de cartographie."),
                  ("Chaîne LiDAR / TF", "Testée après corrections", "Débits observés sans trou et absence d'erreurs temporelles dans l'état propre."),
                  ("Bras et pince", "Actionnement testé", "Mouvements observés après correction du câblage."),
                  ("Saisie autonome complète", "Partiellement validée", "Briques présentes, mais fiabilité de bout en bout insuffisante."),
                  ("Arrêt d'urgence", "Flux logiciel implémenté", "Commande web/MQTT/ROS ; certification matérielle hors périmètre."),
              ], [4.0, 4.1, 8.3], small=True)

    heading(doc, "9.1 Un exemple de validation utile : le temps des scans", 2)
    add_para(doc, "Les deux LiDAR publiaient des mesures correctes, mais le scan fusionné portait un timestamp situé environ une demi-seconde dans le futur. SLAM Toolbox cherchait alors une transformation TF à un instant qui n'existait pas encore. Les scans s'accumulaient dans une file puis étaient rejetés, empêchant la production cohérente de la carte.")
    add_para(doc, "Le nœud `scan_restamper` ne modifie ni les distances ni les angles. Il remplace uniquement le timestamp par l'heure ROS courante moins un décalage configurable de 0,2 seconde, puis republie sur `/scan_fixed`. Le SLAM et les costmaps utilisent ce topic corrigé. Ce petit composant illustre l'intérêt de l'architecture modulaire : une contrainte matérielle a été isolée dans un adaptateur sans contaminer les autres couches.")
    add_callout(doc, "Résultat observé", "Après stabilisation des connexions et correction temporelle, les chaînes `/scan0`, `/scan1`, `/scan_multi` et `/scan_fixed` ont été observées autour de 7,14 Hz, sans trous ni erreurs d'extrapolation dans l'état propre documenté.", fill=PALE, accent=GREEN)

    heading(doc, "9.2 Ce que les essais ont changé dans la conception", 2)
    add_para(doc, "Les essais n'ont pas seulement servi à corriger des anomalies. Ils ont confirmé plusieurs choix : garder les traitements critiques sur le robot, limiter les flux haute fréquence envoyés au web, pouvoir désactiver la caméra pendant la cartographie et conserver des paramètres ajustables sans reconstruire tout le système.")

    page_break(doc)
    heading(doc, "10. Contraintes, risques et limites", 1)
    add_table(doc,
              ["Risque", "Effet possible", "Réponse actuelle", "Reste à faire"],
              [
                  ("Charge du Jetson Nano", "Retards, topics perdus, navigation instable.", "Services limités, fréquences adaptées, caméra arrêtée selon le mode.", "Mesures continues et budgets CPU par mode."),
                  ("Réseau Wi-Fi", "Robot déclaré à tort disponible ou commande retardée.", "Last Will, reconnexion, état de présence.", "Tests de coupure et expiration explicite des commandes."),
                  ("Horloges et timestamps", "Scans et TF incompatibles.", "Restamping des LiDAR et synchronisation de l'hôte.", "Surveillance automatique du décalage."),
                  ("Saisie d'objet", "Mission bloquée ou objet perdu.", "Objet simple, couleur cible, reprise et retour d'erreur.", "Calibration, détection de prise et essais répétés."),
                  ("Sécurité physique", "Mouvement inattendu ou arrêt insuffisant.", "STOP logiciel, règles d'essai, contrôle humain.", "Arrêt matériel, analyse de risques et conformité."),
                  ("Dette documentaire", "Écarts entre code, topics et UML.", "Architecture canonique et contrat MQTT.", "Génération ou vérification automatisée des schémas."),
              ], [3.0, 4.1, 4.9, 4.4], small=True)

    heading(doc, "10.1 Limites du prototype", 2)
    add_para(doc, "RobLaude démontre une architecture et plusieurs fonctions sur un robot réel ; il ne constitue pas encore un produit exploitable sans supervision dans un ERP. La robustesse de la saisie, la gestion des personnes et obstacles dynamiques, la disponibilité réseau et la sûreté de l'arrêt demandent davantage d'essais. La carte et les points doivent aussi être adaptés à chaque bâtiment.")
    add_para(doc, "Le projet n'a pas mesuré le nombre de missions quotidiennes, le coût de maintenance ou le temps réellement économisé par le personnel. Ces données sont indispensables avant une décision de déploiement. La prochaine étape produit doit donc être une expérimentation encadrée, pas une généralisation immédiate.")

    heading(doc, "10.2 Conséquences des arbitrages", 2)
    add_para(doc, "Le découplage améliore la maintenabilité mais multiplie les contrats à surveiller. MQTT absorbe des différences de rythme, mais ajoute un broker et des règles de QoS. ROS 2 accélère la construction de la navigation, mais demande une maîtrise des transformations et du temps. Docker rend l'environnement reproductible, au prix d'une consommation supplémentaire sur une machine limitée. Ces coûts sont acceptables pour le prototype parce qu'ils évitent une réécriture complète de briques complexes.")


def improvements_traceability(doc):
    page_break(doc)
    heading(doc, "11. Axes d'amélioration et évolutivité", 1)
    heading(doc, "11.1 Priorités à court terme", 2)
    add_table(doc,
              ["Priorité", "Action", "Critère de validation"],
              [
                  ("1", "Formaliser un scénario de test de coupure réseau et d'expiration des commandes.", "Aucune commande ancienne exécutée après une reconnexion prolongée."),
                  ("2", "Stabiliser la perception et la saisie sur un objet de référence.", "Série d'essais répétables avec taux de réussite documenté."),
                  ("3", "Ajouter une surveillance des fréquences ROS, de l'horloge et de la charge CPU.", "Alerte avant dégradation de la navigation."),
                  ("4", "Séparer clairement arrêt fonctionnel et arrêt de sécurité matériel.", "Analyse de risques et dispositif physique testé."),
                  ("5", "Aligner automatiquement spécification MQTT, types backend et documentation.", "Détection en CI d'une divergence de contrat."),
              ], [2.0, 8.6, 5.8], small=True)

    heading(doc, "11.2 Évolution multi-robot", 2)
    add_para(doc, "L'arborescence MQTT contient déjà `robotId` et le modèle de données représente plusieurs robots. L'évolution ne consiste donc pas à dupliquer l'application, mais à ajouter une stratégie d'affectation : disponibilité, proximité, batterie et capacités. Le backend resterait l'autorité qui choisit un robot, tandis que chaque bridge ne recevrait que son sous-arbre de commandes.")

    heading(doc, "11.3 Déploiement dans plusieurs bâtiments", 2)
    add_para(doc, "Chaque bâtiment nécessite sa propre carte, ses points et une procédure de validation. Les entités de cartographie déjà présentes constituent une base, mais il faudrait associer clairement utilisateurs, robots, cartes et sites. Les contraintes d'accessibilité et de sécurité devraient être réévaluées pour chaque environnement.")

    heading(doc, "11.4 Pérennité technique", 2)
    add_para(doc, "La pérennité dépend moins du maintien éternel d'une version de React ou ROS 2 que de la stabilité des frontières. Le frontend peut évoluer si l'API reste contractuelle. Le broker peut être remplacé si le contrat de messages est conservé. Un nouveau robot peut être intégré derrière un bridge adapté. Cette indépendance limite le coût des migrations futures.")
    add_callout(doc, "Décision proposée", "Avant d'ajouter des fonctions visibles, investir dans la preuve : tests de reprise réseau, mesures de charge, validation répétée de la saisie et traçabilité automatique des contrats.", fill=PALE_BLUE, accent=NAVY)

    page_break(doc)
    heading(doc, "12. Traçabilité des besoins", 1)
    add_para(doc, "Cette matrice relie le besoin initial à une réponse de conception, à une preuve disponible et à la limite connue. Elle évite de confondre intention, code et validation terrain.")
    add_table(doc,
              ["Besoin", "Réponse de conception", "Preuve", "Limite actuelle"],
              [
                  ("Créer une mission simplement", "PWA responsive et API métier.", "Parcours et code web présents.", "Validation utilisateur PMR à élargir."),
                  ("Connaître l'avancement", "MQTT vers backend puis WebSocket vers PWA.", "Flux temps réel implémenté.", "Dépendance au réseau local."),
                  ("Naviguer dans le bâtiment", "SLAM Toolbox, TF et Nav2.", "Cartographie et navigation réelles documentées.", "Carte à produire pour chaque site."),
                  ("Éviter les obstacles", "LiDAR et costmaps Nav2 sur `/scan_fixed`.", "Chaîne LiDAR stable observée.", "Tests avec public et obstacles dynamiques à renforcer."),
                  ("Arrêter le robot", "Commande dédiée et état de pause.", "Flux logiciel présent.", "Pas une chaîne d'arrêt certifiée."),
                  ("Récupérer un objet", "Caméra, détection, bras et pince.", "Actionnement testé ; briques présentes.", "Fiabilité de bout en bout non démontrée."),
                  ("Administrer plusieurs robots", "Topics par robot et modèle de données.", "Architecture compatible.", "Affectation multi-robot à développer."),
                  ("Maintenir la solution", "Composants séparés et contrats versionnés.", "Packages, tests et documentation.", "Certaines docs historiques ont divergé."),
              ], [3.4, 4.7, 4.4, 3.9], small=True)

    heading(doc, "12.1 Lecture au regard du référentiel", 2)
    add_table(doc,
              ["Compétence", "Éléments apportés"],
              [
                  ("C1.1", "Contexte ERP, acteurs, enjeux métier, objectifs hiérarchisés, retombées attendues et évolutions."),
                  ("C1.2", "Comparaison PWA/native, REST/WebSocket, MQTT/ROS direct, traitement embarqué/déporté ; conséquences exposées."),
                  ("C1.3", "Architecture en composants, contrats MQTT, données, quatre diagrammes, spécifications et évolutivité."),
              ], [2.7, 13.7])


def conclusion_sources(doc):
    page_break(doc)
    heading(doc, "13. Conclusion", 1)
    add_para(doc, "RobLaude s'appuie sur une architecture découpée entre l'interface, le backend, MQTT et le robot. L'utilisateur formule une demande, le backend la valide et la conserve, MQTT transmet les échanges entre les deux environnements, puis le robot exécute la mission et remonte son état. Le fonctionnement interne de ROS 2 reste isolé du système web.")
    add_para(doc, "Les essais sur le robot réel ont validé la navigation, la cartographie et plusieurs flux essentiels. Ils ont aussi montré les limites du matériel et l'importance du temps dans la chaîne LiDAR/TF. La récupération autonome d'objet reste la partie la moins stable du projet et nécessite encore des essais répétés. Cette limite n'annule pas les résultats obtenus, mais elle indique clairement le travail qui reste à faire.")
    add_para(doc, "La prochaine étape consiste à renforcer les tests de sûreté et de reprise réseau, puis à stabiliser la saisie avant d'ajouter de nouvelles fonctions. Le découpage actuel permet ensuite d'envisager plusieurs robots et plusieurs bâtiments sans reprendre toute l'application.")
    add_callout(doc, "Bilan du prototype", "Les fonctions testées sont identifiées, les limites restantes sont connues et les prochaines validations peuvent être planifiées à partir de constats réalisés sur le robot.", fill=PALE, accent=TEAL)

    heading(doc, "Sources et références", 1)
    heading(doc, "Sources internes au projet", 2)
    sources_internal = [
        "[R1] RobLaude, Cahier des charges v1.2, docs/CDC_Roblaude_v1.2.md.",
        "[R2] RobLaude, Architecture et mémoire projet, docs/architecture.md.",
        "[R3] RobLaude, Roadmap v1.4, docs/ROADMAP.md.",
        "[R4] RobLaude, Spécification MQTT, docs/mqtt-spec.md.",
        "[R5] RobLaude, Faits robot sourcés, docs/ROBOT.md.",
        "[R6] RobLaude, Difficultés rencontrées sur le robot, docs/DIFFICULTES_ROBOT.md.",
        "[R7] Dépôt RobLaude, branche feat/demo-soutenance, code frontend, backend et packages ROS 2.",
    ]
    for text in sources_internal:
        add_source(doc, text)

    heading(doc, "Références techniques externes", 2)
    sources_external = [
        "[E1] MDN Web Docs, Making PWAs installable, https://developer.mozilla.org/en-US/docs/Web/Progressive_web_apps/Guides/Making_PWAs_installable, consulté en août 2026.",
        "[E2] OASIS / MQTT.org, MQTT Specification, https://mqtt.org/mqtt-specification/, consulté en août 2026.",
        "[E3] Open Robotics, ROS 2 Humble - Interfaces: topics, services, actions, https://docs.ros.org/en/humble/Concepts/Basic/Interfaces-Topics-Services-Actions.html, consulté en août 2026.",
        "[E4] Nav2, documentation officielle, https://docs.nav2.org/, consultée en août 2026.",
    ]
    for text in sources_external:
        add_source(doc, text)



def build():
    doc = setup_document()
    cover(doc)
    contents_and_summary(doc)
    context_and_needs(doc)
    functional_scope(doc)
    state_of_art(doc)
    architecture(doc)
    contracts_and_data(doc)
    specifications(doc)
    validation_and_limits(doc)
    improvements_traceability(doc)
    conclusion_sources(doc)

    core = doc.core_properties
    core.title = "Dossier de conception B1 - RobLaude"
    core.subject = "Épreuve individuelle Bloc 1 : Conception"
    core.author = "Wissem Karboub"
    core.keywords = "RobLaude, conception, architecture, ROS 2, MQTT, PWA, HETIC"
    core.comments = "Dossier produit à partir des sources du projet RobLaude."

    doc.save(OUTPUT)
    print(f"Document généré : {OUTPUT}")
    print(f"Taille : {OUTPUT.stat().st_size} octets")


if __name__ == "__main__":
    build()
